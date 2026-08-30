from pathlib import Path

import yaml

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


OUT = Path("dist/계약검토_체크리스트_운영방식_임원보고_20260824.docx")
KNOWLEDGE_DIR = Path("knowledge")
FONT = "Apple SD Gothic Neo"
BLACK = "000000"
GRAY = "666666"
LIGHT = "F2F2F2"
MID = "D9D9D9"


def set_run(run, size=10.0, bold=False, color=BLACK):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def fmt_para(p, before=0, after=3.2, line=1.08, keep=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if keep:
        pf.keep_with_next = True


def add_text(doc, text, size=10.0, bold=False, after=3.2, color=BLACK, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    fmt_para(p, after=after)
    set_run(p.add_run(text), size=size, bold=bold, color=color)
    return p


def add_bullet(doc, text, after=2.8):
    p = doc.add_paragraph(style="List Bullet")
    fmt_para(p, after=after, line=1.08)
    p.paragraph_format.left_indent = Mm(4.8)
    p.paragraph_format.first_line_indent = Mm(-3.2)
    set_run(p.add_run(text), size=9.7)
    return p


def add_heading(doc, text):
    p = doc.add_paragraph()
    fmt_para(p, before=4.2, after=2.6, line=1.0, keep=True)
    set_run(p.add_run(text), size=11.5, bold=True)
    return p


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=95, start=115, bottom=95, end=115):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_text(cell, text, bold=False, size=8.8, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    fmt_para(p, after=0, line=1.07)
    set_run(p.add_run(text), size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def set_table_geometry(table, widths_dxa):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "105")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for old in list(grid):
        grid.remove(old)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    tr_pr.append(hdr)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def add_page_field(paragraph):
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rfonts.set(qn("w:eastAsia"), FONT)
    rpr.append(rfonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "15")
    rpr.append(size)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), GRAY)
    rpr.append(color)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = "2"
    run.append(text)
    fld.append(run)
    paragraph._p.append(fld)


def basis_text(check):
    kind = "법령" if check.get("basis") == "statute" else "실무"
    sources = check.get("sources") or []
    if not sources:
        return f"{check.get('severity', '')} / {kind}"
    first = sources[0]
    src = f"{first.get('law', '')} {first.get('article', '')}".strip()
    if len(sources) > 1:
        src += f" 외 {len(sources) - 1}건"
    return f"{check.get('severity', '')} / {src}"


def appendix_header(doc, section):
    section.footer.is_linked_to_previous = False
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fmt_para(p, after=0, line=1.0)
    set_run(p.add_run("별첨 | "), size=7.5, color=GRAY)
    add_page_field(p)


def add_inventory_table(doc, doc_data):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = [1120, 2140, 5100, 1960]
    set_table_geometry(table, widths)
    headers = ["ID", "체크항목", "무엇을 확인하는가", "중요도 / 근거"]
    for i, text in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], text, bold=True, size=7.8, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(table.rows[0].cells[i], LIGHT)
    set_repeat_header(table.rows[0])
    set_row_cant_split(table.rows[0])

    meta = doc_data["meta"]
    module_names = {m["id"]: m["name"] for m in meta.get("modules", [])}
    groups = []
    by_group = {}
    for check in doc_data.get("checks", []):
        group = check.get("module") or "기본·공통"
        if group not in by_group:
            groups.append(group)
            by_group[group] = []
        by_group[group].append(check)

    for group in groups:
        group_row = table.add_row()
        merged = group_row.cells[0].merge(group_row.cells[3])
        set_cell_width(merged, sum(widths))
        group_name = module_names.get(group, group)
        set_cell_text(merged, f"{group_name} ({len(by_group[group])}개)", bold=True, size=7.8)
        shade(merged, "E8E8E8")
        set_row_cant_split(group_row)
        for check in by_group[group]:
            cells = table.add_row().cells
            values = [
                check["id"],
                check.get("label") or check["check"],
                check["check"],
                basis_text(check),
            ]
            for i, value in enumerate(values):
                align = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 3) else WD_ALIGN_PARAGRAPH.LEFT
                set_cell_width(cells[i], widths[i])
                set_cell_text(cells[i], str(value), bold=False, size=7.2, align=align)
                set_cell_margins(cells[i], top=48, start=70, bottom=48, end=70)
            set_row_cant_split(table.rows[-1])
    return table


def add_appendix(doc):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(13.5)
    section.bottom_margin = Mm(13.5)
    section.left_margin = Mm(14)
    section.right_margin = Mm(14)
    section.header_distance = Mm(6)
    section.footer_distance = Mm(7)
    appendix_header(doc, section)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt_para(p, after=2.0, line=1.0)
    set_run(p.add_run("별첨. 계약검토 체크리스트 전체 목록"), size=16.0, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt_para(p, after=5.0, line=1.0)
    set_run(p.add_run("기준일 2026. 8. 24. | 총 240개"), size=8.4, color=GRAY)

    p = doc.add_paragraph()
    fmt_para(p, after=4.0, line=1.08)
    set_run(p.add_run("읽는 방법  "), size=9.0, bold=True)
    set_run(p.add_run("‘체크항목’은 화면의 짧은 항목명, ‘무엇을 확인하는가’는 실제 검토 질문임. 중요도는 필수·권장·참고로, 근거는 법령 조문 또는 계약 실무로 구분함."), size=8.5)

    files = [KNOWLEDGE_DIR / "common.yaml", *sorted((KNOWLEDGE_DIR / "types").glob("*.yaml"))]
    for idx, path in enumerate(files, start=1):
        data = yaml.safe_load(path.read_text(encoding="utf-8"))
        count = len(data.get("checks", []))
        h = doc.add_paragraph()
        fmt_para(h, before=6.0 if idx > 1 else 2.0, after=2.5, line=1.0, keep=True)
        set_run(h.add_run(f"{idx}. {data['meta']['type_name']} ({count}개)"), size=11.0, bold=True)
        add_inventory_table(doc, data)


def add_metric_strip(doc):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = [2580, 2580, 2580, 2580]
    set_table_geometry(table, widths)
    metrics = [
        ("240개", "전체 사전등록 항목"),
        ("169개", "법령·규정 근거"),
        ("71개", "계약 실무 점검"),
        ("129·42·69", "필수·권장·참고"),
    ]
    for i, (number, label) in enumerate(metrics):
        cell = table.rows[0].cells[i]
        shade(cell, LIGHT)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt_para(p, after=0, line=1.0)
        set_run(p.add_run(number + "\n"), size=13.0, bold=True)
        set_run(p.add_run(label), size=8.2, color=GRAY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=125, start=100, bottom=125, end=100)
    return table


def add_flow_table(doc):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    widths = [2580, 2580, 2580, 2580]
    set_table_geometry(table, widths)
    steps = [
        ("1", "항목을 미리 설계", "법무·실무가 질문, 중요도, 적용조건과 근거를 등록"),
        ("2", "계약의 성격을 확인", "표제·파일명·본문에서 유형과 당사 지위를 파악"),
        ("3", "필요한 항목만 선별", "규제 모듈과 개별 조건을 적용하고 관련 원문을 연결"),
        ("4", "사람이 결론을 확정", "적용성·충분성·회사 유불리와 수정 방향을 판단"),
    ]
    for i, (no, title, detail) in enumerate(steps):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt_para(p, after=1.2, line=1.02)
        set_run(p.add_run(no + "\n"), size=12.5, bold=True)
        set_run(p.add_run(title), size=9.6, bold=True)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt_para(p2, after=0, line=1.07)
        set_run(p2.add_run(detail), size=8.3)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=120, start=115, bottom=120, end=115)
        if i in (0, 2):
            shade(cell, LIGHT)
    return table


def add_boundary_table(doc):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [5160, 5160])
    headers = ["도구가 찾아주는 것", "사람이 판단하는 것"]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=9.3, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(table.rows[0].cells[i], LIGHT)
    set_repeat_header(table.rows[0])
    cells = table.add_row().cells
    left = "• 해당 계약에 적용될 가능성이 있는 점검항목\n• 관련 가능성이 있는 계약 조항과 법령 원문\n• 정형 요건의 존재 여부에 관한 보조 표시"
    right = "• 그 항목이 본건에 실제 적용되는지\n• 문언이 충분하고 당사자 간 균형적인지\n• 회사에 불리한지, 수정·협상 또는 위험수용이 필요한지"
    set_cell_width(cells[0], 5160)
    set_cell_width(cells[1], 5160)
    set_cell_text(cells[0], left, size=8.8)
    set_cell_text(cells[1], right, size=8.8)
    return table


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(13.5)
    section.bottom_margin = Mm(13.5)
    section.left_margin = Mm(14)
    section.right_margin = Mm(14)
    section.header_distance = Mm(6)
    section.footer_distance = Mm(6)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.0)
    normal.paragraph_format.space_after = Pt(3.2)
    normal.paragraph_format.line_spacing = 1.08

    list_style = doc.styles["List Bullet"]
    list_style.font.name = FONT
    list_style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    list_style.font.size = Pt(9.7)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt_para(title, after=1.5, line=1.0)
    set_run(title.add_run("계약검토 체크리스트는\n어떻게 만들어지고 사용되는가"), size=18.0, bold=True)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt_para(sub, after=5.0, line=1.0)
    set_run(sub.add_run("질문: 이 항목들은 기계가 만든 것인가, 결국 사람이 확인해야 하는 것인가?  |  2026. 8. 24."), size=8.8, color=GRAY)

    callout = doc.add_table(rows=1, cols=1)
    callout.style = "Table Grid"
    set_table_geometry(callout, [10320])
    shade(callout.cell(0, 0), LIGHT)
    set_cell_margins(callout.cell(0, 0), top=135, start=165, bottom=135, end=165)
    p = callout.cell(0, 0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt_para(p, after=0, line=1.08)
    set_run(p.add_run("답변  "), size=10.7, bold=True)
    set_run(p.add_run("체크항목은 사람이 미리 만들고, 도구는 계약에 맞는 항목과 관련 원문을 찾아줌. 도구가 법적 결론을 내리는 것이 아니며, 최종 판단은 사람이 수행함."), size=10.1, bold=True)

    add_heading(doc, "현재 등록된 체크리스트는 240개임")
    add_metric_strip(doc)
    p = doc.add_paragraph()
    fmt_para(p, before=2.0, after=2.5, line=1.06)
    set_run(p.add_run("구성  "), size=9.4, bold=True)
    set_run(p.add_run("공통 119개 + 계약유형별 121개임. 유형별로는 조달·주주간 각 26, 투자 17, 업무위탁 16, NDA 14, 화해·합의 10, 금융·여신 7, 판매채널 3, 제휴 2개임."), size=8.9)

    add_heading(doc, "체크리스트는 ‘생성’되는 것이 아니라 계약별로 ‘선별’됨")
    add_flow_table(doc)
    p = doc.add_paragraph()
    fmt_para(p, before=2.4, after=2.2, line=1.07)
    set_run(p.add_run("왜 계약마다 다른 항목이 나타나는가  "), size=9.4, bold=True)
    set_run(p.add_run("① 계약 종류, ② 당사의 지위(계약 당사자·수익자), ③ 적용 규제(개인정보·전자금융·클라우드·하도급·담보·계열사 거래), ④ 개별 조건(문서 제목, 공모·사모, 도급·위임, 당사자 역할)을 차례로 적용하기 때문임."), size=8.8)

    p = doc.add_paragraph()
    fmt_para(p, after=2.2, line=1.07)
    set_run(p.add_run("조항 매칭 방식  "), size=9.4, bold=True)
    set_run(p.add_run("계약서를 조항 단위로 나눈 뒤, 항목별 키워드·문장패턴으로 후보를 만들고 문장 유사도(TF-IDF·Jaccard), 조항 제목, 의무·금지 표현, 법령 인용을 함께 점수화함. 점수와 적용조건에 따라 ‘관련 원문 찾음·확인 제안·적용/보완 판단 필요’로 제시하되 법적 결론으로 확정하지 않음."), size=8.7)

    p = doc.add_paragraph()
    fmt_para(p, after=2.2, line=1.07)
    set_run(p.add_run("예시  "), size=9.4, bold=True)
    set_run(p.add_run("‘IT업무위탁계약서’에서 개인정보·정보처리시스템·재위탁 문구가 반복되면 업무위탁 유형과 개인정보·전자금융 모듈이 활성화됨. 이후 ‘재위탁 제한’, ‘목적 외 처리 금지’ 항목과 관련 조항을 연결하며, 실제 개인정보 처리 여부·사전동의 문언의 충분성·수정 필요성은 검토자가 판단함."), size=8.7)

    add_heading(doc, "사람의 검토는 마지막에 덧붙는 절차가 아니라 설계의 핵심임")
    add_boundary_table(doc)
    p = doc.add_paragraph()
    fmt_para(p, before=2.4, after=2.2, line=1.07)
    set_run(p.add_run("표시의 의미  "), size=9.4, bold=True)
    set_run(p.add_run("‘관련 원문 찾음’은 문구가 있다는 뜻일 뿐 ‘문제없음’의 확정이 아님. ‘적용·보완 판단 필요’ 역시 누락을 단정하는 경고가 아니라, 검토자가 계약의 성격과 거래조건을 함께 확인해야 한다는 의미임."), size=8.9)

    conclusion = doc.add_table(rows=1, cols=1)
    conclusion.style = "Table Grid"
    set_table_geometry(conclusion, [10320])
    shade(conclusion.cell(0, 0), LIGHT)
    set_cell_margins(conclusion.cell(0, 0), top=115, start=150, bottom=115, end=150)
    p = conclusion.cell(0, 0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt_para(p, after=0, line=1.05)
    set_run(p.add_run("한 줄 결론  "), size=9.8, bold=True)
    set_run(p.add_run("목적은 사람을 대신한 자동판정이 아니라, 검토 누락을 막고 근거를 찾는 시간을 줄이는 것임."), size=9.5, bold=True)

    p = doc.add_paragraph()
    fmt_para(p, before=2.8, after=0, line=1.0)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run(p.add_run("※ 전체 240개 ID·검토질문·근거는 별첨 참조. 근거: contract-review 지식 YAML·스키마·검증/매칭 로직"), size=7.7, color=GRAY)

    add_appendix(doc)

    doc.core_properties.title = "계약검토 체크리스트는 어떻게 만들어지고 사용되는가"
    doc.core_properties.subject = "사전등록 항목의 계약별 선별과 사람 최종판정 구조"
    doc.core_properties.author = "법무팀"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
